from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).with_name("revised.md")
FIGURE = ROOT / "Knowledge/Docs/proposal/figures/figure_1_verified_architecture.png"
OUT = Path(__file__).with_name("PROPOSAL_SIAGA_FINAL.docx")

NAVY = RGBColor(21, 45, 68)
BLUE = RGBColor(37, 91, 133)
PALE = "EDF3F7"
GRAY = RGBColor(68, 77, 86)


def architecture_figure():
    im = Image.new("RGB", (2400, 960), "white")
    d = ImageDraw.Draw(im)
    font_path = "C:/Windows/Fonts/arial.ttf"
    bold_path = "C:/Windows/Fonts/arialbd.ttf"
    normal = ImageFont.truetype(font_path, 43)
    small = ImageFont.truetype(font_path, 34)
    bold = ImageFont.truetype(bold_path, 47)

    def box(xy, title, detail, fill="#edf3f7"):
        d.rounded_rectangle(xy, radius=22, fill=fill, outline="#587c98", width=4)
        x0, y0, x1, y1 = xy
        d.text(((x0 + x1) / 2, y0 + 26), title, fill="#153350", font=bold, anchor="ma")
        for i, line in enumerate(detail.split("\n")):
            d.text(((x0 + x1) / 2, y0 + 91 + i * 43), line, fill="#30485a", font=small, anchor="ma")

    def arrow(x0, y0, x1, y1):
        d.line((x0, y0, x1, y1), fill="#587c98", width=7)
        d.polygon([(x1, y1), (x1 - 18, y1 - 25), (x1 + 18, y1 - 25)], fill="#587c98")

    box((55, 25, 2345, 170), "Antarmuka pengguna", "Pasien: chat dan asesmen   |   Dokter: supervisi   |   SOC: telemetri", "#e9f3f9")
    arrow(1200, 170, 1200, 220)
    box((55, 220, 2345, 355), "FastAPI dan SIAGA", "Autentikasi, pemeriksaan pesan, keputusan, respons API", "#e9f3f9")
    arrow(1200, 355, 1200, 400)
    layers = [
        (55, 400, 580, 585, "L0", "Normalisasi\nUnicode"),
        (635, 400, 1160, 585, "L1", "Niat dan\nrepresentasi teks"),
        (1215, 400, 1740, 585, "L2", "Sinyal URL dan\nlonjakan pesan"),
        (1795, 400, 2345, 585, "L3 CIM", "Momentum dan\nriwayat sesi"),
    ]
    for x0, y0, x1, y1, title, detail in layers:
        box((x0, y0, x1, y1), title, detail)
    d.text((1200, 632), "Fusi keputusan: ALLOW / WATCH / PROBE / BLOCK", fill="#153350", font=bold, anchor="mm")
    d.line((1200, 658, 1200, 700), fill="#587c98", width=7)
    box((55, 710, 755, 915), "LLM lokal", "Ollama bila aktif;\nrespons fallback", "#eaf5ee")
    box((850, 710, 1550, 915), "DuckDB guardrail", "Hash, fitur, skor;\nTTL state sesi", "#fff5e7")
    box((1645, 710, 2345, 915), "Data aplikasi", "SQLite atau Firestore;\nriwayat chat dan catatan", "#fff0ed")
    FIGURE.parent.mkdir(parents=True, exist_ok=True)
    im.save(FIGURE, dpi=(300, 300))


def set_cell_shading(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "3")
        e.set(qn("w:color"), "D9D9D9")
        borders.append(e)
    tcPr.append(borders)


def set_cell_margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, value in (("top", "90"), ("bottom", "90"), ("left", "95"), ("right", "95")):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), value)
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def add_inline(p, value):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            p.add_run(part.replace("`", ""))


def add_table(doc, rows):
    parsed = [[x.strip() for x in row.strip().strip("|").split("|")] for row in rows]
    parsed = [row for row in parsed if not all(re.fullmatch(r":?-+:?", x or "") for x in row)]
    table = doc.add_table(rows=0, cols=len(parsed[0]))
    table.autofit = False
    usable = 17.0
    widths = [usable / len(parsed[0])] * len(parsed[0])
    if len(widths) == 4:
        widths = [3.1, 4.1, 4.6, 5.2]
    elif len(widths) == 3:
        widths = [3.5, 6.7, 6.8]
    for ridx, row in enumerate(parsed):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cells[cidx].width = Cm(widths[cidx])
            cells[cidx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_borders(cells[cidx])
            set_cell_margins(cells[cidx])
            if ridx == 0:
                set_cell_shading(cells[cidx], "DCE9F2")
            elif ridx % 2 == 0:
                set_cell_shading(cells[cidx], "F8FAFC")
            p = cells[cidx].paragraphs[0]
            p.style = "Table Text"
            add_inline(p, value)
            if ridx == 0:
                for run in p.runs:
                    run.bold = True
        for cell in cells:
            cell._tc.get_or_add_tcPr()
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def main():
    from build_figures import architecture, workflow
    architecture(); workflow()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21)
    sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(0.8); sec.footer_distance = Cm(0.8)
    sec.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor(32, 42, 50)
    normal.paragraph_format.space_after = Pt(6.5)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [
        ("Heading 1", 14, 12, 8),
        ("Heading 2", 10.8, 9, 4),
    ]:
        st = styles[name]
        st.font.name = "Aptos Display" if name == "Heading 1" else "Aptos"
        st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.color.rgb = NAVY
    styles["Title"].font.bold = True
    styles["Subtitle"].font.name = "Aptos"
    styles["Subtitle"].font.size = Pt(13)
    styles["Subtitle"].font.color.rgb = GRAY
    if "Table Text" not in styles:
        from docx.enum.style import WD_STYLE_TYPE
        st = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = styles["Table Text"]
    st.font.name = "Aptos"; st.font.size = Pt(8.6)
    st.paragraph_format.space_after = Pt(1)
    st.paragraph_format.line_spacing = 1.08

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("SIAGA  |  PROPOSAL HACKNUSA 2026").font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    cover = True
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line == "\\pagebreak":
            doc.add_page_break(); cover = False; i += 1; continue
        if line.startswith("# "):
            title = line[2:]
            if cover:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(112)
                add_inline(p, title)
            else:
                p = doc.add_paragraph(title, style="Heading 1")
                if title.startswith("CHAPTER") and title != "CHAPTER 1: INTRODUCTION & BACKGROUND":
                    p.page_break_before = True
        elif line.startswith("## "):
            title = line[3:]
            if cover:
                p = doc.add_paragraph(style="Subtitle")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(24)
                add_inline(p, title)
            else:
                p = doc.add_paragraph(title, style="Heading 2")
                if title == "Referensi":
                    p.page_break_before = True
        elif line.startswith("!["):
            m = re.match(r"!\[.*?\]\((.*?)\)", line)
            if m:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(ROOT / "Knowledge/Docs/proposal" / m.group(1)), width=Cm(16.2))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            add_table(doc, rows)
            continue
        elif line.startswith("**Tabel "):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line)
        elif line.startswith("*Gambar "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, line)
        elif line.startswith("-") and line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif line.startswith("[") and re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(3)
            for part in re.split(r"(https?://\S+)", line):
                if part.startswith("http"):
                    r = p.add_run(part); r.font.size = Pt(8); r.font.color.rgb = BLUE
                else:
                    add_inline(p, part)
        else:
            p = doc.add_paragraph()
            if cover:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(18)
            add_inline(p, line.replace("  ", " "))
        i += 1

    doc.core_properties.title = "Proposal Inovasi SIAGA"
    doc.core_properties.subject = "HackNusa 2026 AI vs AI Defense"
    doc.core_properties.author = "Tim SIAGA"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
