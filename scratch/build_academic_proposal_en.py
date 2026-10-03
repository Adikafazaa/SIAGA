import os
import re
import sys
from pathlib import Path

import docx
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import pypdf
import pypdfium2 as pdfium

ROOT = Path(r"D:\KULIAH-1\ITENAS\HackNusa\Prototype")
BASE_DIR = ROOT / "SIAGA/Knowledge/Docs/proposal"
SOURCE_MD = BASE_DIR / "PROPOSAL_SIAGA_EN_DRAFT.md"
FIGURES_DIR = BASE_DIR / "figures"
DOCX_OUT = BASE_DIR / "PROPOSAL_SIAGA_EN_FINAL.docx"
PDF_OUT = BASE_DIR / "PROPOSAL_SIAGA_EN_FINAL.pdf"

# Palette: strictly academic monochrome / neutral dark slate
DARK = RGBColor(15, 23, 42)      # Deep black-slate for primary text
MUTED = RGBColor(71, 85, 105)    # Slate for metadata/subtitles

def add_runs_with_inline_markdown(p, text, base_font="Times New Roman", base_size=10.0,
                                  base_bold=False, base_italic=False, color=DARK):
    parts = re.split(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        r = p.add_run()
        r.font.name = base_font
        r.font.size = Pt(base_size)
        r.font.color.rgb = color
        if part.startswith("***") and part.endswith("***"):
            r.text = part[3:-3]
            r.bold = True
            r.italic = True
        elif part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]
            r.bold = True
            r.italic = base_italic
        elif part.startswith("*") and part.endswith("*"):
            r.text = part[1:-1]
            r.bold = base_bold
            r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r.text = part[1:-1]
            r.font.name = "Consolas"
            r.font.size = Pt(base_size * 0.90)
            r.bold = base_bold
        else:
            r.text = part
            r.bold = base_bold
            r.italic = base_italic

def set_cell_margins(cell, top=35, bottom=35, left=40, right=40):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(mar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_in_column_table(doc, caption_text, raw_rows, col_widths=None):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(5)
    p_cap.paragraph_format.space_after = Pt(2)
    p_cap.paragraph_format.keep_with_next = True
    add_runs_with_inline_markdown(p_cap, caption_text, base_font="Times New Roman", base_size=8.5, base_bold=True)

    parsed = [[c.strip() for c in r.strip().strip("|").split("|")] for r in raw_rows]
    parsed = [r for r in parsed if not all(re.fullmatch(r":?-+:?", c or "") for c in r)]

    cols_count = len(parsed[0])
    table = doc.add_table(rows=len(parsed), cols=cols_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    usable = 7.8
    if not col_widths or len(col_widths) != cols_count:
        col_widths = [usable / cols_count] * cols_count

    for r_idx, row in enumerate(parsed):
        tr = table.rows[r_idx]._tr
        trPr = tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        cells = table.rows[r_idx].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].width = Cm(col_widths[c_idx])
            cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[c_idx], top=35, bottom=35, left=40, right=40)

            if r_idx == 0:
                tcPr = cells[c_idx]._tc.get_or_add_tcPr()
                b_border = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="0" w:color="334155"/></w:tcBorders>')
                tcPr.append(b_border)

            p = cells[c_idx].paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or (cols_count >= 4 and c_idx >= 1 and len(val) < 15)) else WD_ALIGN_PARAGRAPH.LEFT
            p.alignment = align
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.line_spacing = 1.02
            
            f_size = 7.0 if cols_count >= 5 else 7.5
            add_runs_with_inline_markdown(p, val, base_font="Times New Roman", base_size=f_size,
                                          base_bold=(r_idx == 0))

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(2.0)

def build_proposal_en(line_sp=1.15, p_after=2.8, body_font_sz=10.0):
    doc = Document()

    # Section 0: Cover (1 Column)
    section = doc.sections[0]
    section.page_width = Cm(21.00)
    section.page_height = Cm(29.70)
    section.top_margin = Cm(2.50)
    section.bottom_margin = Cm(2.50)
    section.left_margin = Cm(2.50)
    section.right_margin = Cm(2.50)
    section.different_first_page_header_footer = True

    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_head = p_head.add_run("SIAGA: Progressive AI Dialogue Guardrails for Digital Mental Health | HackNusa 2026")
    r_head.font.name = "Times New Roman"
    r_head.font.size = Pt(8.5)
    r_head.font.italic = True
    r_head.font.color.rgb = MUTED

    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot_space = p_foot.add_run()
    r_foot_space.font.name = "Times New Roman"
    r_foot_space.font.size = Pt(9.0)
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    p_foot._p.append(fldSimple)

    raw_md = SOURCE_MD.read_text(encoding="utf-8")
    lines = raw_md.splitlines()
    i = 0

    # Build Cover Page
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(36)
    p_comp = doc.add_paragraph()
    p_comp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_comp.paragraph_format.space_after = Pt(18)
    r_c1 = p_comp.add_run("HACKNUSA 2026 NATIONAL TECHNOLOGY HACKATHON INNOVATION PROPOSAL\nTOP 30 FINALIST STAGE")
    r_c1.font.name = "Times New Roman"
    r_c1.font.size = Pt(11.0)
    r_c1.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(14)
    r_t = p_title.add_run("SIAGA INNOVATION PROPOSAL")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(18.0)
    r_t.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_s = p_sub.add_run("Progressive AI Conversational Defense for Mental Health Services")
    r_s.font.name = "Times New Roman"
    r_s.font.size = Pt(11.5)
    r_s.italic = True

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(24)
    r_d = p_div.add_run("────────────────────────────────────────────────────────────")
    r_d.font.name = "Times New Roman"
    r_d.font.size = Pt(10.0)
    r_d.font.color.rgb = MUTED

    p_trk = doc.add_paragraph()
    p_trk.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_trk.paragraph_format.space_after = Pt(30)
    r_tk1 = p_trk.add_run("Competition Category / Track:\n")
    r_tk1.font.name = "Times New Roman"
    r_tk1.font.size = Pt(10.0)
    r_tk1.font.color.rgb = MUTED
    r_tk2 = p_trk.add_run("HackNusa 2026 · AI vs AI Defense Track")
    r_tk2.font.name = "Times New Roman"
    r_tk2.font.size = Pt(11.0)
    r_tk2.bold = True

    p_team = doc.add_paragraph()
    p_team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_team.paragraph_format.space_after = Pt(20)
    r_tm1 = p_team.add_run("Proposed By:\n")
    r_tm1.font.name = "Times New Roman"
    r_tm1.font.size = Pt(10.0)
    r_tm1.font.color.rgb = MUTED
    r_tm2 = p_team.add_run("SIAGA Team · Institut Teknologi Nasional Bandung")
    r_tm2.font.name = "Times New Roman"
    r_tm2.font.size = Pt(10.5)
    r_tm2.bold = True

    p_repo = doc.add_paragraph()
    p_repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_repo.paragraph_format.space_after = Pt(36)
    r_rp1 = p_repo.add_run("Open Source Repository (Open Reproducibility):\n")
    r_rp1.font.name = "Times New Roman"
    r_rp1.font.size = Pt(9.5)
    r_rp1.font.color.rgb = MUTED
    r_rp2 = p_repo.add_run("https://github.com/Adikafazaa/SIAGA")
    r_rp2.font.name = "Times New Roman"
    r_rp2.font.size = Pt(9.5)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(36)
    r_dt = p_date.add_run("BANDUNG · OCTOBER 2026")
    r_dt.font.name = "Times New Roman"
    r_dt.font.size = Pt(10.5)
    r_dt.bold = True

    while i < len(lines):
        line = lines[i].strip()
        if line == r"\pagebreak":
            i += 1
            break
        i += 1

    # Section 1: Body (2 Columns - Continuous throughout entire document body)
    sec_body = doc.add_section()
    sec_body.page_width = Cm(21.00)
    sec_body.page_height = Cm(29.70)
    sec_body.top_margin = Cm(2.00)
    sec_body.bottom_margin = Cm(2.00)
    sec_body.left_margin = Cm(2.00)
    sec_body.right_margin = Cm(2.00)
    sec_body.header_distance = Cm(1.20)
    sec_body.footer_distance = Cm(1.20)

    # 2 columns
    cols = parse_xml(r'<w:cols %s w:num="2" w:space="397"/>' % nsdecls('w'))
    sec_body._sectPr.append(cols)

    table_widths = {
        1: [1.7, 1.9, 2.1, 2.1],      # Table 1: 4 cols
        2: [2.1, 1.6, 4.1],           # Table 2: 3 cols
        3: [2.0, 3.2, 2.6],           # Table 3: 3 cols
        4: [2.1, 1.5, 1.4, 1.5, 1.3], # Table 4: 5 cols
        5: [2.3, 1.6, 2.1, 1.8],      # Table 5: 4 cols
        6: [1.8, 2.4, 1.8, 1.8],      # Table 6: 4 cols
        7: [2.0, 2.8, 3.0],           # Table 7: 3 cols
        8: [1.6, 3.4, 2.8],           # Table 8: 3 cols
    }

    fig_w = 7.8
    table_counter = 0

    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True

            r = p.add_run(title.upper())
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.bold = True
            r.font.color.rgb = DARK
            i += 1
            continue

        if line.startswith("## "):
            title = line[3:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True

            r = p.add_run(title)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.0)
            r.bold = True
            r.font.color.rgb = DARK
            i += 1
            continue

        # In-column Images
        if line.startswith("![") or line.startswith("![]"):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if m:
                img_rel = m.group(2)
                img_path = BASE_DIR / img_rel
                if img_path.exists():
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(5)
                    p_img.paragraph_format.space_after = Pt(1.5)
                    p_img.paragraph_format.keep_with_next = True
                    r_img = p_img.add_run()
                    r_img.add_picture(str(img_path), width=Cm(fig_w))

                    # Caption
                    if i + 2 < len(lines) and lines[i+2].strip().startswith("*Figure "):
                        cap_text = lines[i+2].strip()[1:-1]
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(cap_text) > 60 else WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_before = Pt(1)
                        p_cap.paragraph_format.space_after = Pt(4)
                        add_runs_with_inline_markdown(p_cap, cap_text, base_font="Times New Roman", base_size=8.5, base_italic=True)
                        i += 2
            i += 1
            continue

        # Standalone figure caption
        if line.startswith("*Figure ") and line.endswith("*"):
            cap_text = line[1:-1]
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(cap_text) > 60 else WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(1)
            p_cap.paragraph_format.space_after = Pt(4)
            add_runs_with_inline_markdown(p_cap, cap_text, base_font="Times New Roman", base_size=8.5, base_italic=True)
            i += 1
            continue

        # In-column Tables
        if line.startswith("**Table ") and line.endswith("**"):
            cap_text = line
            table_counter += 1
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("|"):
                i += 1
            raw_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw_rows.append(lines[i].strip())
                i += 1

            widths = table_widths.get(table_counter, None)
            add_in_column_table(doc, cap_text, raw_rows, col_widths=widths)
            continue

        # Equation
        if line.startswith("**M_t = clamp(") and line.endswith(").**"):
            eq_text = line[2:-2]
            p_eq = doc.add_paragraph()
            p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_eq.paragraph_format.space_before = Pt(3)
            p_eq.paragraph_format.space_after = Pt(2.5)
            p_eq.paragraph_format.keep_with_next = True
            r_eq = p_eq.add_run(f"{eq_text}  (1)")
            r_eq.font.name = "Times New Roman"
            r_eq.font.size = Pt(8.5)
            r_eq.bold = True
            i += 1
            continue

        # Reference item
        if re.match(r"^\[\d+\]", line):
            p_ref = doc.add_paragraph()
            p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_ref.paragraph_format.left_indent = Cm(0.5)
            p_ref.paragraph_format.first_line_indent = Cm(-0.5)
            p_ref.paragraph_format.space_before = Pt(0)
            p_ref.paragraph_format.space_after = Pt(2.0)
            p_ref.paragraph_format.line_spacing = 1.05
            add_runs_with_inline_markdown(p_ref, line, base_font="Times New Roman", base_size=8.5)
            i += 1
            continue

        # Regular Body Paragraph
        p_body = doc.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body.paragraph_format.space_before = Pt(0)
        p_body.paragraph_format.space_after = Pt(p_after)
        p_body.paragraph_format.line_spacing = line_sp
        p_body.paragraph_format.first_line_indent = Cm(0.4)
        add_runs_with_inline_markdown(p_body, line, base_font="Times New Roman", base_size=body_font_sz)
        i += 1

    doc.save(DOCX_OUT)

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc_obj = word.Documents.Open(str(DOCX_OUT.resolve()))
        doc_obj.SaveAs(str(PDF_OUT.resolve()), FileFormat=17)
        doc_obj.Close()
    finally:
        word.Quit()

    pdf = pdfium.PdfDocument(str(PDF_OUT))
    pages = len(pdf)
    print(f"[EN BUILD SUCCESS] -> Total Pages: {pages}")
    return pages

if __name__ == "__main__":
    build_proposal_en(line_sp=1.15, p_after=2.8, body_font_sz=10.0)
