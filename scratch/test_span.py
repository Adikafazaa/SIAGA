import docx
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import pypdf

doc = Document()
sec1 = doc.sections[0]
sec1.top_margin = Cm(2.0)
sec1.bottom_margin = Cm(2.0)
sec1.left_margin = Cm(2.0)
sec1.right_margin = Cm(2.0)

# Section 1: 2 columns
cols2 = parse_xml(r'<w:cols %s w:num="2" w:space="454"/>' % nsdecls('w'))
sec1._sectPr.append(cols2)

for _ in range(4):
    doc.add_paragraph("Paragraph in 2-column layout. " * 8)

# Continuous break -> 1 column
sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
cols1 = parse_xml(r'<w:cols %s w:num="1"/>' % nsdecls('w'))
sec2._sectPr.append(cols1)

p_wide = doc.add_paragraph("--- FULL WIDTH SPANNING BOTH COLUMNS (FIGURE OR TABLE) ---")
p_wide.runs[0].bold = True

# Continuous break -> 2 columns
sec3 = doc.add_section(WD_SECTION.CONTINUOUS)
cols3 = parse_xml(r'<w:cols %s w:num="2" w:space="454"/>' % nsdecls('w'))
sec3._sectPr.append(cols3)

for _ in range(6):
    doc.add_paragraph("Paragraph back in 2-column layout. " * 8)

test_docx = r"D:\KULIAH-1\ITENAS\HackNusa\Prototype\SIAGA\scratch\test_span.docx"
test_pdf = r"D:\KULIAH-1\ITENAS\HackNusa\Prototype\SIAGA\scratch\test_span.pdf"
doc.save(test_docx)

word = win32com.client.Dispatch("Word.Application")
word.Visible = False
wb = word.Documents.Open(test_docx)
wb.SaveAs(test_pdf, FileFormat=17)
wb.Close()
word.Quit()

reader = pypdf.PdfReader(test_pdf)
print(f"Total pages generated: {len(reader.pages)}")
