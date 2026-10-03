import docx
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import pypdf

doc = Document()

# Section 1: Cover (1 column)
sec1 = doc.sections[0]
sec1.page_width = Cm(21.0)
sec1.page_height = Cm(29.7)
sec1.top_margin = Cm(2.5)
sec1.bottom_margin = Cm(2.5)
sec1.left_margin = Cm(2.5)
sec1.right_margin = Cm(2.5)

p = doc.add_paragraph("PROPOSAL INOVASI SIAGA")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].bold = True
p.runs[0].font.size = Pt(20)

p2 = doc.add_paragraph("This is the cover page in 1 column.")
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Section 2: Body (2 columns)
sec2 = doc.add_section()
sec2.top_margin = Cm(2.0)
sec2.bottom_margin = Cm(2.0)
sec2.left_margin = Cm(2.0)
sec2.right_margin = Cm(2.0)

# Add 2 columns xml
cols = parse_xml(r'<w:cols %s w:num="2" w:space="567"/>' % nsdecls('w')) # 567 twips = 1.0 cm
sec2._sectPr.append(cols)

# Add lots of text to test 2 columns
for ch in range(1, 4):
    h = doc.add_paragraph(f"CHAPTER {ch}: TESTING SECTION")
    h.runs[0].bold = True
    for _ in range(8):
        p = doc.add_paragraph("Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat. Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur.")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

test_docx = r"D:\KULIAH-1\ITENAS\HackNusa\Prototype\SIAGA\scratch\test_twocol.docx"
test_pdf = r"D:\KULIAH-1\ITENAS\HackNusa\Prototype\SIAGA\scratch\test_twocol.pdf"
doc.save(test_docx)

# Convert to PDF
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
wb = word.Documents.Open(test_docx)
wb.SaveAs(test_pdf, FileFormat=17)
wb.Close()
word.Quit()

reader = pypdf.PdfReader(test_pdf)
print(f"Total pages generated: {len(reader.pages)}")
