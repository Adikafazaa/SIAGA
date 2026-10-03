import docx
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import pypdf

doc = Document()

# Section 1: Cover (1-column)
sec1 = doc.sections[0]
sec1.page_width = Cm(21.0)
sec1.page_height = Cm(29.7)
sec1.top_margin = Cm(2.5)
sec1.bottom_margin = Cm(2.5)
sec1.left_margin = Cm(2.5)
sec1.right_margin = Cm(2.5)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(80)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PROPOSAL INOVASI SIAGA")
r.bold = True
r.font.size = Pt(18)
r.font.name = "Times New Roman"

p2 = doc.add_paragraph("Cover Page - 1 Column")
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Section 2: Body (2-column)
sec2 = doc.add_section() # New page for body
sec2.top_margin = Cm(2.0)
sec2.bottom_margin = Cm(2.0)
sec2.left_margin = Cm(2.0)
sec2.right_margin = Cm(2.0)
cols2 = parse_xml(r'<w:cols %s w:num="2" w:space="454"/>' % nsdecls('w'))
sec2._sectPr.append(cols2)

for ch in range(1, 3):
    h = doc.add_paragraph(f"CHAPTER {ch}: TITLE")
    h.runs[0].bold = True
    for _ in range(4):
        p = doc.add_paragraph("Regular 2-column text describing architecture and security principles. " * 6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Wide table (Continuous break -> 1 col)
sec_w1 = doc.add_section(WD_SECTION.CONTINUOUS)
sec_w1._sectPr.append(parse_xml(r'<w:cols %s w:num="1"/>' % nsdecls('w')))

p_tbl = doc.add_paragraph("Table 1. Full Width Table Across Both Columns")
p_tbl.runs[0].bold = True
tbl = doc.add_table(rows=3, cols=4)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in tbl.rows:
    for cell in row.cells:
        cell.width = Cm(4.0)
        cell.paragraphs[0].text = "Cell Content"

# Return to 2 cols (Continuous break -> 2 col)
sec_c1 = doc.add_section(WD_SECTION.CONTINUOUS)
sec_c1._sectPr.append(parse_xml(r'<w:cols %s w:num="2" w:space="454"/>' % nsdecls('w')))

for _ in range(6):
    p = doc.add_paragraph("Paragraph back in 2 columns following the full-width table. " * 6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

out_docx = r"D:\KULIAH-1\ITENAS\HackNusa\Prototype\SIAGA\scratch\test_builder_flow.docx"
out_pdf = r"D:\KULIAH-1\ITENAS\HackNusa\Prototype\SIAGA\scratch\test_builder_flow.pdf"
doc.save(out_docx)

word = win32com.client.Dispatch("Word.Application")
word.Visible = False
wb = word.Documents.Open(out_docx)
wb.SaveAs(out_pdf, FileFormat=17)
wb.Close()
word.Quit()

reader = pypdf.PdfReader(out_pdf)
print(f"Total pages: {len(reader.pages)}")
