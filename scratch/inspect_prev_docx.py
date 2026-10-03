import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

path = r"D:\KULIAH-1\ITENAS\HackNusa\Prototype\SIAGA\Knowledge\Docs\spec\SIAGA_SPESIFIKASI_TEKNIS_ARSITEKTUR_SISTEM (1).docx"
doc = docx.Document(path)
print("Total paragraphs:", len(doc.paragraphs))
print("Total tables:", len(doc.tables))
print("Total sections:", len(doc.sections))

print("\n--- FIRST 30 PARAGRAPHS (COVER & INTRO) ---")
for i, p in enumerate(doc.paragraphs[:30]):
    if p.text.strip():
        runs = [(r.text[:30], r.bold, r.font.name, r.font.size.pt if r.font.size else None, str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None) for r in p.runs]
        print(f"P{i:02d} [{p.alignment}]: '{p.text[:70]}' -> {runs}")

if doc.tables:
    t0 = doc.tables[0]
    print("\nT0 rows:", len(t0.rows), "cols:", len(t0.columns))
    for r in t0.rows[:5]:
        print("  Row:", [c.text.strip()[:30] for c in r.cells])
